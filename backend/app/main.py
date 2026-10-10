from datetime import datetime, timezone, date
import asyncio
import logging
import calendar
from fastapi import Depends, FastAPI, HTTPException, Query, Header, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from sqlalchemy import or_, select, func
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from app.core.config import settings, validate_production_settings
from app.core.security import create_access_token, verify_password
from app.database import Base, engine, get_db
from app.models import Admin, Customer, Product, Purchase, SupportTicket, ServiceReminder, Campaign, AuditLog, MessageLog
from app.schemas import LoginIn, TokenOut, CustomerIn, CustomerOut, ProductIn, ProductOut, PurchaseIn, TicketIn, ReminderIn, CampaignIn, SendMessageIn, CampaignSendIn
from app.integrations.bale.client import BaleClient, BaleAPIError
from app.integrations.bale.bot_service import BaleBotService

validate_production_settings()
Base.metadata.create_all(bind=engine)
logger = logging.getLogger("esbot")
bale_polling_task: asyncio.Task | None = None
app = FastAPI(title="ESBot API", version="0.3.0", description="ESBot CRM و خدمات پس از فروش")
app.add_middleware(CORSMiddleware, allow_origins=[x.strip() for x in settings.cors_origins.split(",") if x.strip()], allow_credentials=True, allow_methods=["GET", "POST", "PATCH", "PUT", "DELETE", "OPTIONS"], allow_headers=["Authorization", "Content-Type", "X-Request-ID"])


@app.on_event("startup")
async def start_bale_bot_polling():
    global bale_polling_task
    if not settings.bale_polling_enabled:
        logger.info("Bale polling disabled by BALE_POLLING_ENABLED")
        return
    if not settings.bale_bot_token or not settings.bale_admin_chat_id:
        logger.warning("Bale bot polling not started: configure BALE_BOT_TOKEN and BALE_ADMIN_CHAT_ID in .env")
        return
    service = BaleBotService(settings.bale_bot_token, settings.bale_admin_chat_id)
    bale_polling_task = asyncio.create_task(service.run_forever(), name="esbot-bale-polling")
    logger.info("Bale bot polling task started")

@app.on_event("shutdown")
async def stop_bale_bot_polling():
    global bale_polling_task
    if bale_polling_task is not None:
        bale_polling_task.cancel()
        try:
            await bale_polling_task
        except asyncio.CancelledError:
            pass
        bale_polling_task = None

@app.get("/")
def root(): return {"name":"ESBot API", "version":"0.3.0", "docs":"/docs", "health":"/api/v1/health"}
@app.get("/api/v1/health")
def health(): return {"status":"ok", "environment":settings.app_env, "database":"connected"}

@app.post("/api/v1/auth/login", response_model=TokenOut)
def login(payload: LoginIn, db: Session = Depends(get_db)):
    admin = db.scalar(select(Admin).where(Admin.username == payload.username))
    if not admin or not admin.is_active or not verify_password(payload.password, admin.password_hash):
        raise HTTPException(401, detail={"code":"INVALID_CREDENTIALS", "message":"نام کاربری یا رمز عبور نادرست است."})
    return TokenOut(access_token=create_access_token(str(admin.id)), username=admin.username)

def current_admin(authorization: str | None = Header(default=None), db: Session = Depends(get_db)):
    import jwt
    if not authorization or not authorization.lower().startswith("bearer "): raise HTTPException(401, "ابتدا وارد سامانه شوید.")
    try:
        payload = jwt.decode(authorization[7:], settings.secret_key, algorithms=["HS256"])
        admin_id = int(payload["sub"])
    except Exception: raise HTTPException(401, "نشست معتبر نیست یا منقضی شده است.")
    admin = db.get(Admin, admin_id)
    if not admin or not admin.is_active: raise HTTPException(401, "حساب کاربری فعال نیست.")
    return admin

def audit(db: Session, admin: Admin, action: str, entity: str, entity_id=None, details=None):
    db.add(AuditLog(admin_id=admin.id, action=action, entity=entity, entity_id=str(entity_id) if entity_id is not None else None, details=details or {}))

def not_found(db, model, ident, message):
    obj=db.get(model, ident)
    if not obj: raise HTTPException(404, message)
    return obj

@app.get("/api/v1/dashboard")
def dashboard(db: Session=Depends(get_db), admin: Admin=Depends(current_admin)):
    return {"customers":db.scalar(select(func.count()).select_from(Customer).where(Customer.is_archived.is_(False))) or 0,
            "products":db.scalar(select(func.count()).select_from(Product).where(Product.is_active.is_(True))) or 0,
            "purchases":db.scalar(select(func.count()).select_from(Purchase)) or 0,
            "open_tickets":db.scalar(select(func.count()).select_from(SupportTicket).where(SupportTicket.status.in_(["open","in_progress"]))) or 0,
            "pending_reminders":db.scalar(select(func.count()).select_from(ServiceReminder).where(ServiceReminder.status=="pending")) or 0,
            "campaigns":db.scalar(select(func.count()).select_from(Campaign)) or 0}

@app.get("/api/v1/customers", response_model=list[CustomerOut])
def list_customers(q: str|None=None, offset:int=Query(0,ge=0), limit:int=Query(50,ge=1,le=100), db:Session=Depends(get_db), admin:Admin=Depends(current_admin)):
    stmt=select(Customer).where(Customer.is_archived.is_(False))
    if q:
        p=f"%{q.strip()}%"; stmt=stmt.where(or_(Customer.first_name.ilike(p),Customer.last_name.ilike(p),Customer.mobile.ilike(p),Customer.email.ilike(p)))
    return list(db.scalars(stmt.order_by(Customer.id.desc()).offset(offset).limit(limit)))
@app.post("/api/v1/customers", response_model=CustomerOut, status_code=201)
def create_customer(payload:CustomerIn, db:Session=Depends(get_db), admin:Admin=Depends(current_admin)):
    item=Customer(**payload.model_dump()); db.add(item)
    try: db.flush(); audit(db,admin,"create","customer",item.id); db.commit(); db.refresh(item)
    except IntegrityError: db.rollback(); raise HTTPException(409,"این شماره موبایل قبلاً ثبت شده است.")
    return item
@app.get("/api/v1/customers/{ident}", response_model=CustomerOut)
def get_customer(ident:int, db:Session=Depends(get_db), admin:Admin=Depends(current_admin)):
    item=not_found(db,Customer,ident,"مشتری پیدا نشد.")
    if item.is_archived: raise HTTPException(404,"مشتری پیدا نشد.")
    return item
@app.patch("/api/v1/customers/{ident}", response_model=CustomerOut)
def update_customer(ident:int,payload:CustomerIn,db:Session=Depends(get_db),admin:Admin=Depends(current_admin)):
    item=not_found(db,Customer,ident,"مشتری پیدا نشد.")
    for k,v in payload.model_dump().items(): setattr(item,k,v)
    try: audit(db,admin,"update","customer",ident); db.commit(); db.refresh(item)
    except IntegrityError: db.rollback(); raise HTTPException(409,"این شماره موبایل قبلاً ثبت شده است.")
    return item
@app.delete("/api/v1/customers/{ident}",status_code=204)
def archive_customer(ident:int,db:Session=Depends(get_db),admin:Admin=Depends(current_admin)):
    item=not_found(db,Customer,ident,"مشتری پیدا نشد."); item.is_archived=True; audit(db,admin,"archive","customer",ident); db.commit()

@app.get("/api/v1/products",response_model=list[ProductOut])
def list_products(q:str|None=None,offset:int=Query(0,ge=0),limit:int=Query(100,ge=1,le=200),db:Session=Depends(get_db),admin:Admin=Depends(current_admin)):
    stmt=select(Product)
    if q: stmt=stmt.where(or_(Product.name.ilike(f"%{q}%"),Product.brand.ilike(f"%{q}%"),Product.model.ilike(f"%{q}%")))
    return list(db.scalars(stmt.order_by(Product.id.desc()).offset(offset).limit(limit)))
@app.post("/api/v1/products",response_model=ProductOut,status_code=201)
def create_product(payload:ProductIn,db:Session=Depends(get_db),admin:Admin=Depends(current_admin)):
    x=Product(**payload.model_dump()); db.add(x); db.flush(); audit(db,admin,"create","product",x.id); db.commit(); db.refresh(x); return x
@app.patch("/api/v1/products/{ident}",response_model=ProductOut)
def update_product(ident:int,payload:ProductIn,db:Session=Depends(get_db),admin:Admin=Depends(current_admin)):
    x=not_found(db,Product,ident,"محصول پیدا نشد.")
    for k,v in payload.model_dump().items(): setattr(x,k,v)
    audit(db,admin,"update","product",ident); db.commit(); db.refresh(x); return x
@app.delete("/api/v1/products/{ident}")
def disable_product(ident:int,db:Session=Depends(get_db),admin:Admin=Depends(current_admin)):
    x=not_found(db,Product,ident,"محصول پیدا نشد."); x.is_active=False; audit(db,admin,"disable","product",ident); db.commit(); return {"ok":True}

@app.get("/api/v1/purchases")
def list_purchases(db:Session=Depends(get_db),admin:Admin=Depends(current_admin)):
    return [{"id":x.id,"customer_id":x.customer_id,"customer_name":f"{x.customer.first_name} {x.customer.last_name}","product_id":x.product_id,"product_name":x.product.name,"purchase_date":x.purchase_date,"invoice_number":x.invoice_number,"service_interval":x.service_interval,"service_interval_unit":x.service_interval_unit,"notes":x.notes} for x in db.scalars(select(Purchase).order_by(Purchase.id.desc()).limit(300))]
@app.post("/api/v1/purchases",status_code=201)
def create_purchase(payload:PurchaseIn,db:Session=Depends(get_db),admin:Admin=Depends(current_admin)):
    not_found(db,Customer,payload.customer_id,"مشتری پیدا نشد."); product=not_found(db,Product,payload.product_id,"محصول پیدا نشد.")
    data=payload.model_dump(); x=Purchase(**data); db.add(x); db.flush()
    # create a due reminder from purchase date and selected interval
    base=date.fromisoformat(x.purchase_date); unit=x.service_interval_unit
    if unit=="days": due=base.fromordinal(base.toordinal()+x.service_interval)
    elif unit=="years":
        year=base.year+x.service_interval; due=date(year,base.month,min(base.day,calendar.monthrange(year,base.month)[1]))
    else:
        idx=base.year*12+base.month-1+x.service_interval; year,month=divmod(idx,12); month+=1
        due=date(year,month,min(base.day,calendar.monthrange(year,month)[1]))
    if product.reminders_enabled: db.add(ServiceReminder(purchase_id=x.id,due_date=due.isoformat()))
    audit(db,admin,"create","purchase",x.id); db.commit(); db.refresh(x); return {"id":x.id,"customer_id":x.customer_id,"product_id":x.product_id,"purchase_date":x.purchase_date,"next_service_date":due.isoformat()}

@app.get("/api/v1/tickets")
def list_tickets(db:Session=Depends(get_db),admin:Admin=Depends(current_admin)):
    return [{"id":x.id,"customer_id":x.customer_id,"customer_name":f"{x.customer.first_name} {x.customer.last_name}","title":x.title,"description":x.description,"priority":x.priority,"status":x.status,"created_at":x.created_at.isoformat()} for x in db.scalars(select(SupportTicket).order_by(SupportTicket.id.desc()).limit(300))]
@app.post("/api/v1/tickets",status_code=201)
def create_ticket(payload:TicketIn,db:Session=Depends(get_db),admin:Admin=Depends(current_admin)):
    not_found(db,Customer,payload.customer_id,"مشتری پیدا نشد."); x=SupportTicket(**payload.model_dump()); db.add(x); db.flush(); audit(db,admin,"create","ticket",x.id); db.commit(); db.refresh(x); return {"id":x.id,"customer_id":x.customer_id,"title":x.title,"description":x.description,"priority":x.priority,"status":x.status,"created_at":x.created_at.isoformat()}
@app.patch("/api/v1/tickets/{ident}")
def update_ticket(ident:int,payload:TicketIn,db:Session=Depends(get_db),admin:Admin=Depends(current_admin)):
    x=not_found(db,SupportTicket,ident,"تیکت پیدا نشد."); not_found(db,Customer,payload.customer_id,"مشتری پیدا نشد.")
    for k,v in payload.model_dump().items(): setattr(x,k,v)
    audit(db,admin,"update","ticket",ident); db.commit(); db.refresh(x); return {"id":x.id,"customer_id":x.customer_id,"title":x.title,"description":x.description,"priority":x.priority,"status":x.status}

@app.get("/api/v1/reminders")
def list_reminders(db:Session=Depends(get_db),admin:Admin=Depends(current_admin)):
    return [{"id":x.id,"purchase_id":x.purchase_id,"due_date":x.due_date,"status":x.status,"message":x.message,"customer_name":f"{x.purchase.customer.first_name} {x.purchase.customer.last_name}","product_name":x.purchase.product.name,"chat_id":x.purchase.customer.bale_chat_id} for x in db.scalars(select(ServiceReminder).order_by(ServiceReminder.due_date).limit(500))]
@app.post("/api/v1/reminders",status_code=201)
def create_reminder(payload:ReminderIn,db:Session=Depends(get_db),admin:Admin=Depends(current_admin)):
    not_found(db,Purchase,payload.purchase_id,"سابقه خرید پیدا نشد."); x=ServiceReminder(**payload.model_dump()); db.add(x); db.flush(); audit(db,admin,"create","reminder",x.id); db.commit(); db.refresh(x); return {"id":x.id,"purchase_id":x.purchase_id,"due_date":x.due_date,"status":x.status,"message":x.message}

@app.get("/api/v1/campaigns")
def list_campaigns(db:Session=Depends(get_db),admin:Admin=Depends(current_admin)):
    return [{"id":x.id,"title":x.title,"message":x.message,"status":x.status,"recipient_count":x.recipient_count,"sent_count":x.sent_count,"created_at":x.created_at.isoformat()} for x in db.scalars(select(Campaign).order_by(Campaign.id.desc()).limit(300))]
@app.post("/api/v1/campaigns",status_code=201)
def create_campaign(payload:CampaignIn,db:Session=Depends(get_db),admin:Admin=Depends(current_admin)):
    x=Campaign(**payload.model_dump()); db.add(x); db.flush(); audit(db,admin,"create","campaign",x.id); db.commit(); db.refresh(x); return {"id":x.id,"title":x.title,"message":x.message,"status":x.status,"recipient_count":x.recipient_count,"sent_count":x.sent_count,"created_at":x.created_at.isoformat()}

@app.post("/api/v1/campaigns/{ident}/send")
async def send_campaign(ident:int, payload:CampaignSendIn, db:Session=Depends(get_db), admin:Admin=Depends(current_admin)):
    campaign=not_found(db,Campaign,ident,"کمپین پیدا نشد.")
    customers=list(db.scalars(select(Customer).where(Customer.is_archived.is_(False),Customer.bale_chat_id.is_not(None))))
    campaign.recipient_count=len(customers); campaign.sent_count=0; campaign.status="sending"; db.commit()
    failed=0
    for customer in customers:
        log=MessageLog(customer_id=customer.id,recipient=str(customer.bale_chat_id),text=campaign.message,status="sending"); db.add(log); db.flush()
        try:
            await BaleClient(payload.token).send_text(str(customer.bale_chat_id),campaign.message)
            log.status="sent"; campaign.sent_count += 1
        except BaleAPIError as exc:
            log.status="failed"; log.error=exc.message; failed += 1
        db.commit()
    campaign.status="sent" if failed==0 else ("partial" if campaign.sent_count else "failed")
    audit(db,admin,"send_campaign","campaign",ident,{"recipients":len(customers),"sent":campaign.sent_count,"failed":failed}); db.commit()
    return {"campaign_id":ident,"recipient_count":len(customers),"sent_count":campaign.sent_count,"failed_count":failed,"status":campaign.status}

@app.post("/api/v1/bale/test")
async def bale_test(payload:SendMessageIn,admin:Admin=Depends(current_admin)):
    try:
        client=BaleClient(payload.token); me=await client.get_me()
        return {"ok":True,"bot":me}
    except BaleAPIError as exc: raise HTTPException(502,detail={"code":exc.code,"message":exc.message})
@app.post("/api/v1/bale/send")
async def bale_send(payload:SendMessageIn,db:Session=Depends(get_db),admin:Admin=Depends(current_admin)):
    log=MessageLog(recipient=payload.chat_id,text=payload.text,status="sending"); db.add(log); db.flush()
    try:
        result=await BaleClient(payload.token).send_text(payload.chat_id,payload.text)
        log.status="sent"; audit(db,admin,"send","bale_message",log.id,{"recipient":payload.chat_id}); db.commit()
        return {"ok":True,"message_id":result.get("message_id"),"log_id":log.id}
    except BaleAPIError as exc:
        log.status="failed"; log.error=exc.message; db.commit(); raise HTTPException(502,detail={"code":exc.code,"message":exc.message})
@app.get("/api/v1/messages")
def list_messages(db:Session=Depends(get_db),admin:Admin=Depends(current_admin)):
    return [{"id":x.id,"recipient":x.recipient,"text":x.text,"status":x.status,"error":x.error,"created_at":x.created_at.isoformat()} for x in db.scalars(select(MessageLog).order_by(MessageLog.id.desc()).limit(200))]
@app.get("/api/v1/audit")
def list_audit(db:Session=Depends(get_db),admin:Admin=Depends(current_admin)):
    return [{"id":x.id,"admin_id":x.admin_id,"action":x.action,"entity":x.entity,"entity_id":x.entity_id,"details":x.details,"created_at":x.created_at.isoformat()} for x in db.scalars(select(AuditLog).order_by(AuditLog.id.desc()).limit(200))]

@app.post("/api/v1/customers/import")
def import_customers(file:UploadFile=File(...),db:Session=Depends(get_db),admin:Admin=Depends(current_admin)):
    if not file.filename or not file.filename.lower().endswith(".xlsx"): raise HTTPException(400,"فقط فایل Excel با پسوند xlsx پشتیبانی می‌شود.")
    from openpyxl import load_workbook
    import io
    try: wb=load_workbook(io.BytesIO(file.file.read()),read_only=True,data_only=True); ws=wb.active
    except Exception: raise HTTPException(400,"فایل Excel معتبر نیست.")
    rows=ws.iter_rows(values_only=True); header=next(rows,None)
    if not header: raise HTTPException(400,"فایل خالی است.")
    keys=[str(v or '').strip().lower() for v in header]; required={"first_name","last_name","mobile"}
    if not required.issubset(set(keys)): raise HTTPException(400,"ستون‌های first_name، last_name و mobile الزامی هستند.")
    inserted=0; skipped=0
    for row in rows:
        d={keys[i]:row[i] for i in range(min(len(keys),len(row)))}; mobile=str(d.get("mobile") or "").strip()
        if not mobile or db.scalar(select(Customer).where(Customer.mobile==mobile)): skipped+=1; continue
        db.add(Customer(first_name=str(d.get("first_name") or "").strip(),last_name=str(d.get("last_name") or "").strip(),mobile=mobile,email=str(d.get("email") or "").strip() or None,notes=str(d.get("notes") or "").strip() or None)); inserted+=1
    audit(db,admin,"import","customer",details={"inserted":inserted,"skipped":skipped}); db.commit(); return {"inserted":inserted,"skipped":skipped}
@app.get("/api/v1/customers/export")
def export_customers(db:Session=Depends(get_db),admin:Admin=Depends(current_admin)):
    from openpyxl import Workbook
    import io
    wb=Workbook(); ws=wb.active; ws.title="Customers"; ws.append(["first_name","last_name","mobile","email","bale_user_id","bale_chat_id","notes"])
    for x in db.scalars(select(Customer).where(Customer.is_archived.is_(False)).order_by(Customer.id)):
        ws.append([x.first_name,x.last_name,x.mobile,x.email,x.bale_user_id,x.bale_chat_id,x.notes])
    stream=io.BytesIO(); wb.save(stream); stream.seek(0)
    return StreamingResponse(stream,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":"attachment; filename=esbot-customers.xlsx"})
